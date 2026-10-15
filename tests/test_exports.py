import csv
import io
from decimal import Decimal
import pytest
from openpyxl import load_workbook
from django.core.management import call_command
from core.models import Role, Job, Customer, Policy
from core.jobs import safe_cell
from test_workflows import issue_pickup, post

@pytest.mark.parametrize('value', ['=SUM(A1:A2)', '+cmd', '-2+3', '@NOW()','\t=bad'])
def test_formula_cells_are_inert(value):
    assert safe_cell(value).startswith("'")

@pytest.mark.django_db(transaction=True)
def test_hq_export_csv_and_xlsx_preserve_filters_and_redact_formulas(data,client_for,settings,tmp_path):
    settings.PRIVATE_ROOT=tmp_path
    bill=issue_pickup(data)
    customer=Customer.objects.get(pk=bill.customer_id);customer.name='=HYPERLINK("https://bad")';customer.save()
    hq=client_for(data['users'][Role.HQ])
    payload={'feature':'customers','format':'csv','filters':{'franchise':str(data['books'][0].pk)}}
    first=post(hq,'exports/',payload,'export-csv')
    assert first.status_code==202,first.data
    assert post(hq,'exports/',payload,'export-csv').data==first.data
    assert Job.objects.count()==1
    call_command('run_jobs',once=True)
    job=Job.objects.get();assert job.state=='done'
    response=hq.get(f'/api/v1/jobs/{job.pk}/download/')
    rows=list(csv.reader(io.StringIO(b''.join(response.streaming_content).decode('utf-8-sig'))))
    assert rows[0][0]=='id' and any(cell.startswith("'=HYPERLINK") for row in rows for cell in row)
    xlsx=post(hq,'exports/',{'feature':'customers','format':'xlsx','filters':{}},'export-xlsx')
    assert xlsx.status_code==202
    call_command('run_jobs',once=True)
    xjob=Job.objects.order_by('-created_at').first();book=load_workbook(tmp_path/xjob.file_name,data_only=False)
    assert book.active['A1'].value=='id'
    assert all(cell.data_type=='s' for row in book.active.iter_rows() for cell in row if cell.value is not None)

@pytest.mark.django_db
def test_exports_require_hq_and_current_scope(data,client_for):
    bill=issue_pickup(data)
    payload={'feature':'bills','format':'csv','filters':{}}
    assert post(client_for(data['users'][Role.APP]),'exports/',payload).status_code==403
    assert post(client_for(data['users'][Role.WAREHOUSE]),'exports/',payload).status_code==403
    Policy.objects.filter(key='export_delivery').delete()
    assert post(client_for(data['users'][Role.HQ]),'exports/',payload).status_code==409

@pytest.mark.django_db(transaction=True)
def test_revoked_export_job_cannot_deliver(data,client_for,settings,tmp_path):
    settings.PRIVATE_ROOT=tmp_path
    issue_pickup(data);hq=data['users'][Role.HQ];client=client_for(hq)
    result=post(client,'exports/',{'feature':'bills','format':'csv','filters':{}},'revoked-export')
    assert result.status_code==202
    hq.permission_version+=1;hq.save()
    call_command('run_jobs',once=True)
    job=Job.objects.get();assert job.state=='failed'
    assert client_for(hq).get(f'/api/v1/jobs/{job.pk}/download/').status_code==404
