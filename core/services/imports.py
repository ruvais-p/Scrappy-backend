import csv
import io
import hashlib
from datetime import datetime,date,time
from decimal import Decimal,InvalidOperation
from itertools import islice
from zipfile import BadZipFile
from django.db import connection,transaction
from openpyxl import load_workbook
from django.utils import timezone
from rest_framework.exceptions import ValidationError
from core.models import *
from core.policy import require,policy,PolicyUnavailable
from .common import event,Conflict
COLUMNS=['customer_name','phone','address','area','upi_id','franchise_code','driver_email','pickup_date','pickup_time','items','latitude','longitude']

def _cell(value):
    """Return a stable text representation while preserving phone strings."""
    if value is None:
        return ''
    if isinstance(value,datetime):
        return value.isoformat(sep=' ')
    if isinstance(value,(date,time)):
        return value.isoformat()
    return str(value).strip()

def _parse_date(value):
    if isinstance(value,datetime):
        return value.date()
    if isinstance(value,date):
        return value
    return datetime.strptime(str(value).strip(),'%d-%m-%Y').date()

def _parse_time(value):
    if isinstance(value,datetime):
        return value.time().replace(second=0,microsecond=0)
    if isinstance(value,time):
        return value.replace(second=0,microsecond=0)
    raw=str(value).strip()
    for pattern in ('%I:%M %p','%H:%M:%S','%H:%M'):
        try:
            return datetime.strptime(raw,pattern).time()
        except ValueError:
            continue
    raise ValueError()

def preview(actor,upload):
    require(actor,'imports')
    if upload.size>5*1024*1024: raise ValidationError('File limit is 5 MB.')
    raw=upload.read()
    if upload.name.lower().endswith('.csv'):
        try: reader=csv.reader(io.StringIO(raw.decode('utf-8-sig')));headers=next(reader);records=list(islice(reader,5001))
        except (UnicodeError,csv.Error,StopIteration): raise ValidationError('Invalid UTF-8 CSV.')
    elif upload.name.lower().endswith('.xlsx'):
        import zipfile
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                if sum(i.file_size for i in z.infolist())>30*1024*1024: raise ValidationError('Workbook expanded size is too large.')
            wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
            try:
                if 'Leads' not in wb.sheetnames: raise ValidationError('Workbook must contain a Leads sheet.')
                rows=wb['Leads'].iter_rows(values_only=True);headers=list(next(rows));records=list(islice(rows,5001))
            finally: wb.close()
        except (BadZipFile,KeyError,ValueError,StopIteration): raise ValidationError('Invalid XLSX workbook.')
    else: raise ValidationError('Upload CSV or XLSX only.')
    if headers!=COLUMNS or len(records)>5000: raise ValidationError('Use the exact template columns and at most 5,000 rows.')
    results=[]
    for n,row in enumerate(records,2):
        values=list(row)
        d={k:_cell(values[i]) if i<len(values) else '' for i,k in enumerate(COLUMNS)};errors=[]
        if len(values)!=len(COLUMNS) or any(len(_cell(v))>2000 for v in values): errors.append('Invalid column count or cell length.')
        if any(_cell(v).lstrip().startswith('=') for v in values): errors.append('Formulas are not permitted.')
        if not d.get('customer_name') or not d.get('phone'): errors.append('Customer name and phone required for preview.')
        if not 7<=len(''.join(c for c in d.get('phone','') if c.isdigit()))<=15: errors.append('Phone must contain 7 to 15 digits.')
        f=Franchise.objects.filter(code=d.get('franchise_code'),kind='scrap',status='active').first()
        if not f: errors.append('Unknown active scrap franchise.')
        driver=User.objects.filter(email=d.get('driver_email'),franchise=f,is_active=True,role__in=[Role.DRIVER,Role.COLLECTOR]).first() if d.get('driver_email') else None
        if d.get('driver_email') and not driver: errors.append('Driver must be active in the specified franchise.')
        elif driver: errors.append('Do not assign a collector in the import. HQ routes each lead to its franchise admin first.')
        try:
            pickup_date=_parse_date(values[7] if len(values)>7 else d.get('pickup_date',''))
            if pickup_date<timezone.localdate(): raise ValueError()
            slot=_parse_time(values[8] if len(values)>8 else d['pickup_time']) if d.get('pickup_time') else time(9)
            d['pickup_date']=pickup_date.isoformat();d['pickup_time']=slot.isoformat()
        except ValueError: errors.append('Use a date today or later in DD-MM-YYYY and time in hh:mm AM/PM.')
        try:
            lat,lon=d.get('latitude',''),d.get('longitude','')
            if bool(lat)!=bool(lon): raise ValueError()
            if lat:
                lat_n,lon_n=Decimal(lat),Decimal(lon)
                if (not lat_n.is_finite() or not lon_n.is_finite() or
                    not -90<=lat_n<=90 or not -180<=lon_n<=180): raise ValueError()
        except (InvalidOperation,ValueError,TypeError): errors.append('Enter a valid coordinate pair.')
        d['franchise']=str(f.id) if f else None;d['driver']=driver.id if driver else None
        results.append({'row':n,'data':d,'errors':errors,'lead':None})
    return ImportBatch.objects.create(actor=actor,digest=hashlib.sha256(raw).hexdigest(),rows=results)

@transaction.atomic
def commit(actor,batch):
    require(actor,'imports');config=policy('lead_import')
    locked_batch=ImportBatch.objects.select_for_update().get(pk=batch.pk)
    # Keep an already-held model instance coherent for service callers that
    # render the commit response immediately after this function returns.
    batch.rows=locked_batch.rows;batch.committed=locked_batch.committed;batch.version=locked_batch.version
    batch.actor_id=locked_batch.actor_id;batch.digest=locked_batch.digest
    if batch.actor_id!=actor.pk and actor.role not in (Role.SUPER,Role.ADMIN,Role.HQ):
        raise Conflict('This import is outside your permitted scope.')
    if config.get('commit')!='whole_file' or config.get('duplicates')!='reject_same_file': raise PolicyUnavailable('This import policy is not supported.')
    if batch.committed: return batch
    if any(r['errors'] for r in batch.rows): raise ValidationError('Correct all invalid rows before importing.')
    # Serialize identical files even when different HQ accounts submit them.
    with connection.cursor() as cursor:
        cursor.execute('SELECT pg_advisory_xact_lock(%s)',[int(batch.digest[:16],16) % (2**63)])
    if ImportBatch.objects.filter(digest=batch.digest,committed=True).exclude(pk=batch.pk).exists(): raise Conflict('This file has already been imported.')
    for r in batch.rows:
        d=r['data'];f=Franchise.objects.get(pk=d['franchise'],status='active')
        if d['driver']: raise Conflict('Direct collector assignment is no longer supported. Preview the import again without a driver.')
        if datetime.strptime(d['pickup_date'],'%Y-%m-%d').date()<timezone.localdate(): raise Conflict('Pickup date has passed. Preview again.')
        phone=''.join(c for c in d['phone'] if c.isdigit())
        if not 7<=len(phone)<=15: raise ValidationError('Invalid phone.')
        customer,_=Customer.objects.get_or_create(franchise=f,phone=phone,defaults={'name':d['customer_name'],'address':d['address'],'area':d['area']})
        lead=Lead.objects.create(franchise=f,customer=customer,state='unassigned',pickup_date=d['pickup_date'],pickup_time=d['pickup_time'],items=d['items'],upi_id=d['upi_id'],latitude=d['latitude'] or None,longitude=d['longitude'] or None)
        r['lead']=str(lead.pk);r['status']='created';r['message']='Created; waiting for HQ to assign a franchise admin.'
        event(actor,'leads',lead,action='lead.imported')
    batch.committed=True;batch.save(update_fields=['rows','committed','version','updated_at'])
    return batch
