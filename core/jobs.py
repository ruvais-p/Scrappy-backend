import csv
import io
import json
import uuid
import urllib.request
import urllib.error
from urllib.parse import urlparse
from datetime import timedelta
from pathlib import Path
from openpyxl import Workbook
from django.conf import settings
from django.db import transaction
from django.http import FileResponse
from django.utils import timezone
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.exceptions import PermissionDenied,NotFound,ValidationError
from drf_spectacular.utils import extend_schema
from .models import Job,User,Role,AuditEvent
from .policy import policy,require
from .views import feature_scope,get_object
from .serializers import FIELDS
from .serializers import FIELDS,serialize
from .services.common import command

EXPORTS={'bills','customers','payments','earnings','audit','inventory','rates','receipts'}

def redact_audit(value):
    blocked=('password','secret','token','authorization','cookie','utr','reference')
    if isinstance(value,dict): return {str(k):redact_audit(v) for k,v in value.items() if not any(word in str(k).lower() for word in blocked)}
    if isinstance(value,list): return [redact_audit(v) for v in value]
    return value

def audit_delivery_payload(event):
    return {'event_id':str(event.pk),'action':event.action,'resource':event.resource,'franchise':str(event.franchise_id) if event.franchise_id else None,'actor':str(event.actor_id) if event.actor_id else None,'created_at':event.created_at.isoformat(),'details':redact_audit(event.details)}
def export_rows(actor,payload):
    feature=payload['feature']
    if feature not in EXPORTS: raise ValidationError('Unsupported export dataset.')
    if actor.role not in (Role.SUPER,Role.HQ): raise PermissionDenied('Bulk export is limited to Super Admin and HQ Manager.')
    cfg=policy('export_delivery')
    if cfg.get('boundary')!='all_filtered': raise ValidationError('Configure the export boundary before continuing.')
    if actor.role in (Role.SUPER,Role.HQ):
        # Export authority is deliberately narrower than screen visibility, but
        # both authorized roles receive a network-scoped snapshot.
        qs=FIELDS[feature][0].objects.all()
    else:
        require(actor,feature)
        qs=feature_scope(actor,feature)
    from .queries import filter_records
    filters=payload.get('filters',{})
    allowed={'search','from','to','issued','franchise','state','kind','product','warehouse','driver','receipt','requirement','flagged','settlement_state','action','actor','franchise_type','staff','customer','direction','financial_only','payment_state','category','status','active','method'}
    if not isinstance(filters,dict) or set(filters)-allowed: raise ValidationError('Unsupported export filter.')
    qs,_=filter_records(qs,feature,filters)
    if qs.count()>10000: raise ValidationError('Narrow this export to 10,000 records or fewer.')
    return serialize(feature,qs,True)

def safe_cell(value):
    text=json.dumps(value) if isinstance(value,(dict,list)) else str(value if value is not None else '')
    if text.lstrip().startswith(('=','+','-','@','\t','\r','\n')):return "'"+text
    return text

def process(job):
    actor=User.objects.get(pk=job.actor_id,is_active=True)
    if job.kind=='account_invitation':
        from .auth import issue_token
        from .views import TeamView
        require(actor,'team');target=User.objects.get(pk=job.payload['user'],is_active=True)
        if target.role not in TeamView().targets(actor) or (actor.franchise_id and actor.franchise_id!=target.franchise_id): raise PermissionDenied()
        issue_token(target,job.payload['purpose']);return
    if actor.permission_version!=job.payload.get('permission_version'):raise PermissionDenied('Permissions changed; request the document again.')
    if job.kind=='audit_delivery':
        config=policy('audit_delivery')
        if not config.get('enabled') or not config.get('endpoint'): raise ValidationError('Audit delivery policy is not enabled.')
        endpoint=urlparse(str(config['endpoint']))
        if endpoint.scheme!='https' or not endpoint.netloc: raise ValidationError('Audit delivery endpoint must use HTTPS.')
        event=AuditEvent.objects.get(pk=job.payload['event'])
        payload=json.dumps(audit_delivery_payload(event),separators=(',',':')).encode()
        request=urllib.request.Request(endpoint.geturl(),data=payload,headers={'Content-Type':'application/json','User-Agent':'scrappy-audit/1'},method='POST')
        with urllib.request.urlopen(request,timeout=min(int(config.get('timeout_seconds',5)),10)) as response:
            if response.status<200 or response.status>=300: raise RuntimeError(f'audit delivery returned {response.status}')
        job.expires_at=timezone.now()+timedelta(days=30)
        return
    if job.kind=='export':
        rows=export_rows(actor,job.payload);columns=list(rows[0]) if rows else ['No records'];kind=job.payload['format']
        if kind=='csv':
            text=io.StringIO();writer=csv.writer(text);writer.writerow(columns)
            for row in rows:writer.writerow([safe_cell(row.get(k)) for k in columns])
            data=text.getvalue().encode('utf-8-sig')
        else:
            wb=Workbook();ws=wb.active;ws.title='Records';ws.append(columns)
            for row in rows:
                ws.append([safe_cell(row.get(k)) for k in columns])
                for cell in ws[ws.max_row]:cell.data_type='s'
            b=io.BytesIO();wb.save(b);data=b.getvalue()
        name=str(uuid.uuid4())+'.'+kind
    elif job.kind=='pdf':
        from .documents import render_current_bill
        bill=get_object(actor,job.payload['feature'],job.payload['bill'])
        if bill.state=='draft':raise ValidationError('Draft has no issued PDF.')
        data,rendition_version=render_current_bill(bill.pk)
        job.payload={**job.payload,'rendition_version':rendition_version}
        name=str(uuid.uuid4())+'.pdf'
    else:raise ValidationError('Unknown job type.')
    settings.PRIVATE_ROOT.mkdir(mode=0o700,parents=True,exist_ok=True)
    path=settings.PRIVATE_ROOT/name;path.write_bytes(data);path.chmod(0o600)
    job.file_name=name;job.expires_at=timezone.now()+timedelta(hours=24)

class ExportView(APIView):
    @extend_schema(request=dict,responses=dict)
    def post(self,request):
        feature=request.data.get('feature');kind=request.data.get('format')
        if feature not in EXPORTS or kind not in ('csv','xlsx'):raise ValidationError('Choose a supported export format and dataset.')
        payload={'feature':feature,'format':kind,'filters':request.data.get('filters',{}),'permission_version':request.user.permission_version}
        export_rows(request.user,payload)
        def run(actor):
            export_rows(actor,payload)
            j=Job.objects.create(actor=actor,kind='export',payload=payload)
            return serialize('jobs',j)
        return Response(command(request.user,request.headers.get('Idempotency-Key'),'export.create',payload,run),status=202)

class AuditDeliveryView(APIView):
    @extend_schema(request=dict,responses=dict)
    def post(self,request):
        if request.user.role!=Role.SUPER: raise PermissionDenied('Only Super Admin can configure audit delivery.')
        config=policy('audit_delivery')
        if not config.get('enabled') or not config.get('endpoint'): raise ValidationError('Audit delivery policy is not enabled.')
        endpoint=urlparse(str(config['endpoint']))
        if endpoint.scheme!='https' or not endpoint.netloc: raise ValidationError('Audit delivery endpoint must use HTTPS.')
        event=AuditEvent.objects.filter(pk=request.data.get('event')).first()
        if not event: raise NotFound()
        payload={'event':str(event.pk),'permission_version':request.user.permission_version}
        def enqueue(actor):
            job=Job.objects.create(actor=actor,kind='audit_delivery',payload=payload,expires_at=timezone.now()+timedelta(days=30))
            return serialize('jobs',job)
        return Response(command(request.user,request.headers.get('Idempotency-Key'),'audit-delivery.create',payload,enqueue),status=202)

class JobStatusView(APIView):
    @extend_schema(responses=dict)
    def get(self,request,pk):
        job=Job.objects.filter(pk=pk,actor=request.user).first()
        if not job: raise NotFound()
        return Response(serialize('jobs',job))

class DownloadView(APIView):
    @extend_schema(responses=bytes)
    def get(self,request,pk):
        job=Job.objects.filter(pk=pk,actor=request.user,state='done',expires_at__gt=timezone.now()).first()
        if not job or not job.file_name:raise NotFound()
        if request.user.permission_version!=job.payload.get('permission_version'):raise PermissionDenied()
        if job.kind=='export':export_rows(request.user,job.payload)
        else:
            bill=get_object(request.user,job.payload['feature'],job.payload['bill'])
            if job.payload.get('rendition_version')!=bill.version:
                from .documents import render_current_bill,pdf_response
                data,version=render_current_bill(bill.pk)
                return pdf_response(data,version,'scrappy-pdf.pdf')
        path=settings.PRIVATE_ROOT/job.file_name
        if path.parent!=settings.PRIVATE_ROOT or not path.exists():raise NotFound()
        response=FileResponse(path.open('rb'),as_attachment=True,filename='scrappy-'+job.kind+path.suffix)
        response['Cache-Control']='private, no-store'
        response['X-Content-Type-Options']='nosniff'
        if job.kind=='pdf': response['X-Document-Version']=str(job.payload['rendition_version'])
        return response
